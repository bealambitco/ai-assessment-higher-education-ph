"""Build Analysis_Workbook_Extension.xlsx for the extension arms (re-runnable, read-only on sources).

Matches the house style of Analysis_Workbook_FINAL.xlsx and Analysis_Workbook_Supplement_v2.xlsx: Arial,
navy header band, banded rows, frozen header, every count beside its denominator, a source note under every
table. Primary-study files are not read or modified; primary study figures quoted here come from the re-run of
the checks over the answers already collected.

No raw answer text is written to the workbook, and no model-to-answer mapping beyond the labels that are
already public in the extension results files.

  python3 -B code/extensions/round2/build_extension_workbook.py \
      --package ~/.../30_ROUND2_EXTENSION \
      --out     ~/.../30_ROUND2_EXTENSION/analysis/Analysis_Workbook_Extension.xlsx
"""
import argparse
import csv
import json
import sys
from datetime import datetime
from collections import Counter, defaultdict
from pathlib import Path

sys.dont_write_bytecode = True

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.data_source import AxDataSource, StrRef
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Font as DFont, Paragraph, ParagraphProperties, RegularTextRun
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ----------------------------------------------------------------------------------------------
# 0. ARGUMENTS AND CONSTANTS
# ----------------------------------------------------------------------------------------------
ARMS = ["access", "documents", "discovery"]
ARM_LABEL = {"access": "Low-resource models (arm G)", "documents": "Whole documents (arm C)",
             "discovery": "Rule discovery (arm D)"}
MODEL_LABEL = {
    "gemma3-12b": "Gemma 3 12B", "gptoss20b": "GPT-OSS 20B", "qwen3-30b-a3b": "Qwen3 30B A3B",
    "gemini31pro": "Gemini 3.1 Pro Preview", "kimik3": "Kimi K3",
}
PRIMARY_LABEL = {"astra": "Astra (OpenAI Codex)", "luna": "Luna (OpenAI Codex)",
                 "fable": "Fable 5.1 (Claude)", "haiku": "Haiku 4.5 (Claude)"}
PRIMARY_ORDER = ["astra", "luna", "fable", "haiku"]
OUTCOME_LABEL = {"acceptable": "Scored acceptable", "serious": "Scored serious (Major or Critical)",
                 "confirmed_error": "Confirmed wrong by hand", "unscored": "Not scored by a person"}
OUTCOME_ORDER = ["acceptable", "serious", "confirmed_error", "unscored"]
RULE_LABEL = {"v1": "v1 (frozen 15 Sep 2026)", "v2": "v2 (revised)", "v2b": "v2b (exploratory variant)"}
NOT_HERE = "not in these files"


def parse_args(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--package", required=True, help="the 30_ROUND2_EXTENSION package directory")
    ap.add_argument("--out", required=True, help="workbook to write")
    ap.add_argument("--billing", default=None,
                    help="OpenRouter activity export (default: openrouter_activity_2026-09-20.csv beside the package)")
    return ap.parse_args(argv)


# ----------------------------------------------------------------------------------------------
# 1. LOAD SOURCES
# ----------------------------------------------------------------------------------------------
def read_csv(path):
    with Path(path).open(newline="") as f:
        return list(csv.DictReader(f))


def as_bool(v):
    return str(v).strip().lower() == "true"


def as_float(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def as_int(v):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


def median(xs):
    xs = sorted(x for x in xs if x is not None)
    if not xs:
        return None
    n = len(xs)
    return xs[n // 2] if n % 2 else round((xs[n // 2 - 1] + xs[n // 2]) / 2, 3)


def load(pkg, billing_path):
    """Read every source once. Returns a plain dict; nothing here writes."""
    pkg = Path(pkg).expanduser().resolve()
    src = {
        "per_run": {a: read_csv(pkg / f"analysis/{a}_per_run.csv") for a in ARMS},
        "answers": read_csv(pkg / "analysis/checks_v2_per_answer.csv"),
        "changes": read_csv(pkg / "analysis/checks_v2_changes.csv"),
        "package": pkg,
    }
    # per-request provenance: one kept record per run (the latest attempt)
    records = {}
    for p in sorted(pkg.glob("work/*/by_model/*/run_records/*.json")):
        r = json.loads(p.read_text())
        records[(r["arm"], r["label"], r["run_id"])] = r
    src["records"] = records
    # superseded replies kept beside the reply that replaced them (deviation D2-01)
    superseded = Counter()
    for p in sorted(pkg.glob("work/*/by_model/*/raw_outputs/*_raw_superseded_*.txt")):
        superseded[(p.parents[3].name, p.parents[1].name)] += 1
    src["superseded_files"] = superseded
    # scheduled runs per arm, from the manifest fixed before collection
    src["scheduled"] = {a: len(json.loads((pkg / f"work/{a}/MANIFEST.json").read_text())["runs"]) for a in ARMS}
    src["cases"] = {a: json.loads((pkg / f"work/{a}/MANIFEST.json").read_text())["cases"] for a in ARMS}
    bill = Path(billing_path).expanduser() if billing_path else pkg.parent / "openrouter_activity_2026-09-20.csv"
    src["billing"] = read_csv(bill)
    src["billing_name"] = bill.name
    run_record = pkg / "08_RESULTS_PUBLIC/checks_v2_run_record.json"
    src["checks_run_record"] = json.loads(run_record.read_text()) if run_record.exists() else {}
    return src


# ----------------------------------------------------------------------------------------------
# 2. MEASURES
# ----------------------------------------------------------------------------------------------
def arm_rows(src, arm):
    """One measure row per configuration in an arm, in a fixed order."""
    out = []
    for label in sorted({r["label"] for r in src["per_run"][arm]}):
        rows = [r for r in src["per_run"][arm] if r["label"] == label]
        recs = [src["records"][(arm, label, r["run_id"])] for r in rows if (arm, label, r["run_id"]) in src["records"]]
        quoted = [r for r in rows if as_bool(r["parsed"]) and as_int(r["citations"])]
        d1 = Counter(r["decision_v1"] for r in rows)
        d2 = Counter(r["decision_v2"] for r in rows)
        out.append({
            "label": label,
            "name": MODEL_LABEL.get(label, label),
            "model_id": rows[0]["model_id"],
            "scheduled": src["scheduled"][arm],
            "runs": len(rows),
            "records": len(recs),
            "replies_saved": sum(1 for r in recs if (r.get("reply_chars") or 0) > 0),
            "transport_errors": sum(1 for r in rows if r["error"]),
            "parsed": sum(1 for r in rows if as_bool(r["parsed"])),
            "unparseable": sum(1 for r in rows if not as_bool(r["parsed"])),
            "schema_ok": sum(1 for r in rows if as_bool(r["schema_ok"])),
            "format_failures": sum(1 for r in rows if not as_bool(r["schema_ok"])),
            "answers_with_quotes": len(quoted),
            "quotes": sum(as_int(r["citations"]) or 0 for r in quoted),
            "quotes_in_evidence": sum(as_int(r["citations_in_evidence"]) or 0 for r in quoted),
            "quotes_in_excerpt": sum(as_int(r["citations_in_excerpt"]) or 0 for r in quoted),
            "released_v1": d1["release_with_warning"], "routed_v1": d1["route"], "blocked_v1": d1["block"],
            "released_v2": d2["release_with_warning"], "routed_v2": d2["route"], "blocked_v2": d2["block"],
            "consequential": sum(1 for r in rows if r["proposed_action"] in
                                 ("record_grade", "approve_request", "penalize_student")),
            "median_seconds": median([as_float(r["seconds"]) for r in rows]),
            "cost": round(sum(as_float(r["cost_usd"]) or 0.0 for r in rows), 6),
            "prompt_tokens": sum(as_int(r["prompt_tokens"]) or 0 for r in rows),
            "completion_tokens": sum(as_int(r["completion_tokens"]) or 0 for r in rows),
            "truncated_kept": sum(1 for r in recs if r.get("finish_reason") == "length"),
            "retried": sum(1 for r in recs if (r.get("attempts") or 1) > 1),
            "raised_cap": sum(1 for r in recs if (r.get("max_tokens") or 0) > 4000),
            "superseded_files": src["superseded_files"].get((arm, label), 0),
        })
    return out


def primary_rows(src):
    """Primary-study configurations, re-run under both rule sets in arm A."""
    out = []
    for m in PRIMARY_ORDER:
        rows = [r for r in src["answers"] if r["model"] == m]
        rec = [r for r in rows if r["recorded_v1_decision"]]
        item = {"model": m, "name": PRIMARY_LABEL[m], "answers": len(rows),
                "recorded": len(rec),
                "reproduced": sum(1 for r in rec if r["recorded_v1_decision"] == r["decision_v1"]),
                "format_failures": sum(1 for r in rows if "schema" in (r["failed_checks_v1"] or "").split(";"))}
        for v in ("v1", "v2", "v2b"):
            c = Counter(r["decision_" + v] for r in rows)
            item[f"released_{v}"], item[f"routed_{v}"], item[f"blocked_{v}"] = \
                c["release_with_warning"], c["route"], c["block"]
        out.append(item)
    return out


def effect_rows(src, rules):
    """Effect of a rule set by outcome class, over all 192 re-run answers."""
    out = []
    for oc in OUTCOME_ORDER:
        sub = [r for r in src["answers"] if r["outcome_class"] == oc]
        out.append({
            "outcome": oc, "name": OUTCOME_LABEL[oc], "answers": len(sub),
            "withheld_v1": sum(1 for r in sub if as_bool(r["withheld_v1"])),
            "withheld_new": sum(1 for r in sub if as_bool(r["withheld_" + rules])),
            "released_by_new": sum(1 for r in sub if as_bool(r["withheld_v1"]) and not as_bool(r["withheld_" + rules])),
            "withheld_by_new": sum(1 for r in sub if not as_bool(r["withheld_v1"]) and as_bool(r["withheld_" + rules])),
        })
    return out


def change_reason(row):
    """Plain-language reason, derived only from the columns in checks_v2_changes.csv."""
    new = [c for c in (row["checks_new"] or "").split(";") if c]
    old = [c for c in (row["checks_v1"] or "").split(";") if c]
    modes = {m.split("=")[0]: m.split("=")[1] for m in (row["match_modes"] or "").split(";") if "=" in m}
    if row["direction"] == "recovered":
        if old == ["quoted_source_membership"]:
            return ("Quotation found in the supplied evidence once curly quotation marks, dashes and spacing "
                    "were normalized (rule R3).")
        if all(c.startswith("recompute_") for c in old) and old:
            return ("Arithmetic check passed under the revised comparison (R1 unit-label normalization or R4 "
                    "rounding tolerance); the expected quantity was matched by name.")
        return "Withheld under the frozen checks, released under the revised ones."
    if new == ["consequential_action_routing"]:
        return ("A decision phrase in the answer text triggered routing, although the declared action was not "
                "consequential (v2b rule B1).")
    if any(c.startswith("recompute_") for c in new):
        mode = sorted({modes.get(c) for c in new if modes.get(c)})
        if mode == ["ambiguous_candidates"]:
            return ("Expected quantity not reported under its own name and several unit-compatible candidates "
                    "existed, so the arithmetic check failed closed (rule R2b).")
        if mode == ["matched_in_prose"]:
            return ("Expected quantity absent from the structured results; a value found in the answer text was "
                    "compared and did not match (v2b rule B2).")
        return "Arithmetic check failed under the revised rules."
    return "Released under the frozen checks, withheld under the revised ones."


def cost_rows(src):
    """Billed against kept, per model, from the provider's own export and the run records."""
    kept = defaultdict(lambda: {"runs": 0, "cost": 0.0, "prompt": 0, "completion": 0, "arms": set()})
    kept_gids = set()
    for (arm, label, _rid), r in src["records"].items():
        mid = r["model_id"]
        k = kept[mid]
        k["runs"] += 1
        k["cost"] += float((r.get("usage") or {}).get("cost") or 0.0)
        k["prompt"] += int((r.get("usage") or {}).get("prompt_tokens") or 0)
        k["completion"] += int((r.get("usage") or {}).get("completion_tokens") or 0)
        k["arms"].add(arm)
        if r.get("generation_id"):
            kept_gids.add(r["generation_id"])
    billed = defaultdict(lambda: {"calls": 0, "cost": 0.0, "prompt": 0, "completion": 0, "reasoning": 0})
    kept_billed = defaultdict(float)
    superseded = defaultdict(lambda: {"calls": 0, "cost": 0.0, "truncated": 0, "no_content": 0})
    for b in src["billing"]:
        # the export carries a dated permaslug; the run records carry the routing id it was called with
        slug = b["model_permaslug"]
        bl = billed[slug]
        bl["calls"] += 1
        bl["cost"] += float(b["cost_total"] or 0)
        bl["prompt"] += int(b["tokens_prompt"] or 0)
        bl["completion"] += int(b["tokens_completion"] or 0)
        bl["reasoning"] += int(b["tokens_reasoning"] or 0)
        if b["generation_id"] in kept_gids:
            kept_billed[slug] += float(b["cost_total"] or 0)
        else:
            s = superseded[slug]
            s["calls"] += 1
            s["cost"] += float(b["cost_total"] or 0)
            s["truncated"] += 1 if b["finish_reason_normalized"] == "length" else 0
            s["no_content"] += 1 if int(b["tokens_reasoning"] or 0) >= int(b["tokens_completion"] or 0) > 0 else 0
    # join billing slugs to run-record model ids on the leading path, which is stable across dated variants
    def base(slug):
        return slug.rsplit("-2026", 1)[0]
    rows = []
    for mid in sorted(kept):
        slug = next((s for s in billed if base(s) == mid), None)
        b = billed.get(slug, {"calls": 0, "cost": 0.0, "prompt": 0, "completion": 0, "reasoning": 0})
        s = superseded.get(slug, {"calls": 0, "cost": 0.0, "truncated": 0, "no_content": 0})
        rows.append({"model_id": mid, "slug": slug, "arms": sorted(kept[mid]["arms"]),
                     "billed_calls": b["calls"], "billed_cost": b["cost"],
                     "kept_runs": kept[mid]["runs"], "kept_cost": kept[mid]["cost"],
                     "kept_billed_cost": kept_billed.get(slug, 0.0),
                     "superseded_calls": s["calls"], "superseded_cost": s["cost"],
                     "superseded_truncated": s["truncated"], "superseded_no_content": s["no_content"],
                     "kept_prompt": kept[mid]["prompt"], "kept_completion": kept[mid]["completion"],
                     "billed_prompt": b["prompt"], "billed_completion": b["completion"],
                     "billed_reasoning": b["reasoning"]})
    unmatched = sorted(set(billed) - {r["slug"] for r in rows})
    totals = {
        "billed_calls": len(src["billing"]),
        "billed_cost": sum(float(b["cost_total"] or 0) for b in src["billing"]),
        "kept_runs": len(src["records"]),
        "kept_cost": sum(float((r.get("usage") or {}).get("cost") or 0.0) for r in src["records"].values()),
        "kept_gids": len(kept_gids),
        "gids_missing_from_export": len(kept_gids - {b["generation_id"] for b in src["billing"]}),
        "billed_prompt": sum(int(b["tokens_prompt"] or 0) for b in src["billing"]),
        "billed_completion": sum(int(b["tokens_completion"] or 0) for b in src["billing"]),
        "billed_reasoning": sum(int(b["tokens_reasoning"] or 0) for b in src["billing"]),
        "kept_prompt": sum(int((r.get("usage") or {}).get("prompt_tokens") or 0) for r in src["records"].values()),
        "kept_completion": sum(int((r.get("usage") or {}).get("completion_tokens") or 0) for r in src["records"].values()),
        "unmatched_slugs": unmatched,
    }
    totals["superseded_calls"] = totals["billed_calls"] - totals["kept_runs"]
    totals["kept_billed_cost"] = sum(kept_billed.values())
    totals["superseded_cost"] = sum(s["cost"] for s in superseded.values())
    # the export and the run records report the same requests at different precision
    totals["cost_reporting_gap"] = totals["kept_cost"] - totals["kept_billed_cost"]
    return rows, totals


def collected_on(src):
    """Latest collection date in the retained run records (not the build date), so the workbook is stable."""
    days = sorted({(r.get("finished_at") or "")[:10] for r in src["records"].values() if r.get("finished_at")})
    return days[-1] if days else NOT_HERE


# ----------------------------------------------------------------------------------------------
# 3. STYLE HELPERS  (as in _build/build_27_workbook.py, so the three workbooks sit together)
# ----------------------------------------------------------------------------------------------
NAVY, BLUE, RED, GOLD, GREY = "0B1F3A", "0038A8", "CE1126", "8A6100", "8A94A6"
DARKGREY, INK, MUTED = "4A5363", "1F2937", "5B6576"
FONT = "Arial"
HFILL = PatternFill("solid", fgColor=NAVY)
BAND = PatternFill("solid", fgColor="F3F5F8")
TILEFILL = PatternFill("solid", fgColor="F3F5F8")
BOT = Border(bottom=Side(style="thin", color="D5DAE1"))
F_BODY = Font(name=FONT, size=10, color=INK)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=16, bold=True, color=NAVY)
F_SUB = Font(name=FONT, size=10, italic=True, color=MUTED)
F_H2 = Font(name=FONT, size=12, bold=True, color=BLUE)
F_NOTE = Font(name=FONT, size=8, italic=True, color=MUTED)
F_BOLD = Font(name=FONT, size=10, bold=True, color=INK)
USD = '"US$"#,##0.00'
USD4 = '"US$"#,##0.0000'
INT = "#,##0"
DEC1 = "#,##0.0"


def sheet(wb, name, title, subtitle, widths=None, grid=False):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = grid
    ws["A1"] = title
    ws["A1"].font = F_TITLE
    ws["A2"] = subtitle
    ws["A2"].font = F_SUB
    ws["A2"].alignment = Alignment(vertical="top")
    ws.row_dimensions[1].height = 24
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A4"
    return ws


def h2(ws, row, text, col=1):
    ws.cell(row=row, column=col, value=text).font = F_H2
    return row + 1


def note(ws, row, text, col=1):
    c = ws.cell(row=row, column=col, value=text)
    c.font = F_NOTE
    c.alignment = Alignment(wrap_text=False, vertical="top")
    return row + 1


def table(ws, row, headers, rows, fmts=None, col=1, wrap_header=True, bold_first=False):
    """Write a styled table. Returns (first_data_row, last_data_row, next_free_row)."""
    for j, h in enumerate(headers):
        c = ws.cell(row=row, column=col + j, value=h)
        c.font = F_HEAD
        c.fill = HFILL
        c.alignment = Alignment(wrap_text=wrap_header, vertical="center", horizontal="left" if j == 0 else "center")
    ws.row_dimensions[row].height = 30 if wrap_header else 18
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            c = ws.cell(row=row + 1 + i, column=col + j, value=v)
            c.font = F_BOLD if (bold_first and j == 0) else F_BODY
            c.border = BOT
            if i % 2 == 1:
                c.fill = BAND
            if fmts and j < len(fmts) and fmts[j]:
                c.number_format = fmts[j]
            left = j == 0 or (isinstance(v, str) and not str(v).startswith("="))
            c.alignment = Alignment(horizontal="left" if left else "center", vertical="top",
                                    wrap_text=isinstance(v, str) and len(v) > 40)
    return row + 1, row + len(rows), row + len(rows) + 2


def cp(sz=900, bold=False, color=INK):
    return CharacterProperties(sz=sz, b=bold, solidFill=color, latin=DFont(typeface=FONT))


def rich(sz=900, bold=False, color=INK):
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp(sz, bold, color)), endParaRPr=cp(sz, bold, color))])


def mk_title(text, sz=1100, bold=True, color=NAVY):
    para = Paragraph(pPr=ParagraphProperties(defRPr=cp(sz, bold, color)), r=[RegularTextRun(rPr=cp(sz, bold, color), t=text)])
    return Title(tx=Text(rich=RichText(p=[para])), overlay=False)


def bar(ws, cat_ref, data_ref, title, y_title=None, colors=None, horizontal=False, y_max=None,
        numfmt="0", width=17, height=8.5, gap=60, legend=True):
    ch = BarChart()
    ch.type = "bar" if horizontal else "col"
    ch.grouping = "clustered"
    ch.gapWidth = gap
    ch.add_data(data_ref, titles_from_data=True)
    ch.set_categories(cat_ref)
    for s in ch.series:
        s.cat = AxDataSource(strRef=StrRef(f=s.cat.numRef.f))
    ch.title = mk_title(title)
    ch.style = 2
    ch.width, ch.height = width, height
    palette = colors or [BLUE]
    for i, s in enumerate(ch.series):
        colr = palette[i % len(palette)]
        s.graphicalProperties = GraphicalProperties(solidFill=colr)
        s.graphicalProperties.line = LineProperties(solidFill=colr)
        s.dLbls = DataLabelList()
        s.dLbls.showVal = True
        for a in ("showSerName", "showCatName", "showLegendKey", "showPercent", "showBubbleSize"):
            setattr(s.dLbls, a, False)
        s.dLbls.numFmt = numfmt
        s.dLbls.txPr = rich(800, True, INK)
        s.dLbls.position = "outEnd"
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.majorGridlines = None
    ch.y_axis.numFmt = numfmt
    ch.y_axis.scaling.min = 0
    if y_max is not None:
        ch.y_axis.scaling.max = y_max
    ch.x_axis.txPr = rich(900, False, INK)
    ch.y_axis.txPr = rich(800, False, MUTED)
    ch.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
    if y_title:
        ch.y_axis.title = mk_title(y_title, 900, False, MUTED)
    if horizontal:
        ch.x_axis.scaling.orientation = "maxMin"
        ch.y_axis.crosses = "max"
    if legend:
        ch.legend.position = "b"
        ch.legend.txPr = rich(900, False, INK)
    else:
        ch.legend = None
    return ch


def place(ws, ch, anchor, source_text):
    import math
    import re
    ws.add_chart(ch, anchor)
    m = re.match(r"([A-Z]+)(\d+)", anchor)
    colr, row = m.group(1), int(m.group(2))
    rows = math.ceil(ch.height / 0.53) + 1
    c = ws[f"{colr}{row + rows}"]
    c.value = "Source: " + source_text
    c.font = F_NOTE
    return row + rows + 2


def ref(ws, c1, r1, c2, r2):
    return Reference(ws, min_col=c1, min_row=r1, max_col=c2, max_row=r2)


def chart_block(ws, row, headers, rows):
    """A visible block holding the exact values a chart plots, as on the primary study sheets."""
    for j, h in enumerate(headers):
        c = ws.cell(row=row, column=1 + j, value=h)
        c.font = F_BODY
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            ws.cell(row=row + 1 + i, column=1 + j, value=v).font = F_BODY
    return row


# ----------------------------------------------------------------------------------------------
# 4. SHEETS
# ----------------------------------------------------------------------------------------------
UNSCORED_LINE = ("Correctness is not scored for the three collection arms: no person has scored a "
                 "low-resource, whole-document or rule-discovery answer. Everything on these sheets "
                 "describes what the models produced and what the checks did with it.")


def build(src, out_path):
    wb = Workbook()
    wb.remove(wb.active)
    arms = {a: arm_rows(src, a) for a in ARMS}
    prim = primary_rows(src)
    crows, ctot = cost_rows(src)
    day = collected_on(src)
    rr = src["checks_run_record"]

    ws_readme = sheet(wb, "README", "Extension: analysis workbook",
                      "Three new collection arms and one re-analysis, added beside the closed primary study. "
                      "Built by code/extensions/round2/build_extension_workbook.py from the extension package; "
                      f"answers collected {day}.", [30, 120])
    ws_dash = sheet(wb, "Dashboard", "What the extension adds",
                    "Headline counts, each beside its denominator. Correctness for the new arms is not yet scored "
                    "by a person; see the note under the tiles.",
                    [2, 13, 13, 13, 2, 13, 13, 13, 2, 13, 13, 13, 2, 13, 13, 13])
    ws_dash.freeze_panes = None

    # ---------------------------------------------------------------- Revised checks
    ws = sheet(wb, "Revised checks", "Arm A: what the revised checks change",
               "The frozen checks (v1), the revised checks (v2) and the exploratory variant (v2b), re-run over the "
               "192 answers already collected in the primary study. No new answers, no paid calls, no score revised.",
               [26, 12, 12, 12, 12, 12, 12, 12, 13, 13, 13, 13])
    r = h2(ws, 4, "Table 1. Release decisions by configuration and rule set (48 answers per configuration; 192 in all)")
    rows = []
    for p in prim:
        for v in ("v1", "v2", "v2b"):
            rows.append([p["name"], RULE_LABEL[v], p["answers"], p[f"released_{v}"], p[f"routed_{v}"],
                         p[f"blocked_{v}"], p[f"released_{v}"] + p[f"routed_{v}"] + p[f"blocked_{v}"]])
    _, _, r = table(ws, r, ["Configuration", "Rule set", "Answers", "Released with a warning", "Routed to a person",
                            "Blocked", "Decisions recorded"], rows,
                    [None, None, INT, INT, INT, INT, INT], bold_first=True)
    r = note(ws, r - 1, "Source: analysis/checks_v2_per_answer.csv (192 rows, 48 per configuration). Every answer "
                        "receives exactly one decision under each rule set, so the last column equals the answer count.")
    r += 1
    r = h2(ws, r, "Table 2. Reproduction of the decisions recorded at the September 18 score lock")
    rows = [[p["name"], p["answers"], p["recorded"], p["reproduced"],
             p["recorded"] - p["reproduced"], p["format_failures"]] for p in prim]
    rows.append(["All four configurations", sum(p["answers"] for p in prim), sum(p["recorded"] for p in prim),
                 sum(p["reproduced"] for p in prim),
                 sum(p["recorded"] - p["reproduced"] for p in prim), sum(p["format_failures"] for p in prim)])
    _, _, r = table(ws, r, ["Configuration", "Answers re-run", "Decisions on record from the primary study",
                            "Reproduced exactly", "Mismatches", "Format-check failures under v1"], rows,
                    [None, INT, INT, INT, INT, INT], bold_first=True)
    r = note(ws, r - 1, "Only the 96 primary answers (Astra, Luna) carry a decision recorded at the score lock; the 96 "
                        "Claude answers were checked after it, so there is nothing to reproduce for them.")
    r += 1
    r = h2(ws, r, "Table 3. Effect by outcome class, over all 192 answers")
    rows = []
    for v in ("v2", "v2b"):
        for e in effect_rows(src, v):
            rows.append([RULE_LABEL[v], e["name"], e["answers"], e["withheld_v1"], e["withheld_new"],
                         e["released_by_new"], e["withheld_by_new"]])
    _, _, r = table(ws, r, ["Rule set", "Outcome class", "Answers", "Withheld under v1", "Withheld under the new rules",
                            "Released by the new rules", "Withheld by the new rules"], rows,
                    [None, None, INT, INT, INT, INT, INT], bold_first=True)
    r = note(ws, r - 1, "Withheld = routed to a person or blocked. Outcome classes come from the primary study human scores, "
                        "unchanged: 48 answers scored acceptable, 2 confirmed wrong by hand, 0 scored serious, 142 not "
                        "scored. An unscored answer is counted neither as recovered help nor as contained error.")
    if rr.get("pre_specified_reading"):
        psr = rr["pre_specified_reading"]
        r = note(ws, r, "Pre-specified reading, fixed in CHECKS_V2_SPEC.md before the run: recovered acceptable "
                        f"{psr['recovered_acceptable']}, newly contained confirmed errors {psr['newly_contained_errors']}, "
                        f"new withholding of acceptable answers {psr['new_withholding_of_acceptable']}; verdict "
                        f"'{psr['verdict']}'.")
    r += 1
    r = h2(ws, r, "Table 4. Every answer whose decision changed, with the reason the checks give")
    chg = sorted(src["changes"], key=lambda x: (x["rules"], PRIMARY_ORDER.index(x["model"]), x["case_id"], x["run_id"]))
    rows = [[c["case_id"], PRIMARY_LABEL[c["model"]], c["run_id"].rsplit("_", 1)[-1], RULE_LABEL[c["rules"]],
             OUTCOME_LABEL[c["outcome_class"]], c["from"].replace("release_with_warning", "released with a warning")
             .replace("block", "blocked").replace("route", "routed"),
             c["to"].replace("release_with_warning", "released with a warning").replace("block", "blocked")
             .replace("route", "routed"),
             "Recovered" if c["direction"] == "recovered" else "Newly withheld",
             c["checks_v1"] or "none", c["checks_new"] or "none", change_reason(c)] for c in chg]
    _, _, r = table(ws, r, ["Case", "Configuration", "Repetition", "Rule set", "Outcome class", "Decision under v1",
                            "Decision under the new rules", "Direction", "Checks failed under v1",
                            "Checks failed under the new rules", "Why the decision moved"], rows,
                    [None] * 11, bold_first=True)
    ws.column_dimensions["K"].width = 74
    r = note(ws, r - 1, f"{len(rows)} changed decisions out of 384 answer-by-rule-set comparisons (192 answers under v2 "
                        "and under v2b). Source: analysis/checks_v2_changes.csv. Reasons are derived from the check "
                        "names and match modes in that file; no answer text is reproduced here.")
    r += 1
    r = h2(ws, r, "Chart data (values used by the chart below)")
    cd0 = chart_block(ws, r, ["Configuration", "Blocked under v1", "Blocked under v2"],
                      [[p["name"], p["blocked_v1"], p["blocked_v2"]] for p in prim])
    r += len(prim) + 2
    ch = bar(ws, ref(ws, 1, cd0 + 1, 1, cd0 + len(prim)), ref(ws, 2, cd0, 3, cd0 + len(prim)),
             "The revised checks block fewer answers in three of four configurations (of 48 each)",
             "Answers blocked (of 48)", [GREY, BLUE], y_max=12, width=18)
    r = place(ws, ch, f"A{r}", "analysis/checks_v2_per_answer.csv (Table 1 above).")
    note(ws, r, "Blocking is not by itself a good or a bad outcome: whether a blocked answer was wrong is a scoring "
                "question, and 142 of the 192 answers are still unscored.")

    # ---------------------------------------------------------------- the three collection arms
    arm_sheet(wb, src, "access", arms["access"], prim,
              "Low-resource models", "Arm G: open-weight models a school could self-host",
              "Three open-weight models in the 8B-30B range answered the same 24 cases, with the same supplied "
              "excerpt and the same required JSON format as the primary study, through a hosted aggregator.")
    arm_sheet(wb, src, "documents", arms["documents"], None,
              "Whole documents", "Arm C: the whole retained document instead of an excerpt",
              "Two models answered all 24 cases with the entire governing document attached, so the model had to "
              "find the passage itself. Quotations are matched against that document, not against the excerpt.")
    arm_sheet(wb, src, "discovery", arms["discovery"], None,
              "Rule discovery", "Arm D: no passage pre-selected",
              "Two models answered eight cases with every retained document of that institution attached and no "
              "indication of which passage governs. The answer key and the checks are unchanged.")

    # ---------------------------------------------------------------- Costs and provenance
    ws = sheet(wb, "Costs and provenance", "Costs and provenance of the extension calls",
               f"The provider's own export ({src['billing_name']}) against the retained run records. Billed covers "
               "every request charged; kept covers the latest attempt of each run, which is what the analysis uses.",
               [34, 13, 14, 12, 14, 14, 15, 15, 15, 16, 16, 16])
    r = h2(ws, 4, "Table 1. Billed against kept, by model")
    rows = [[c["model_id"], ", ".join(ARM_LABEL[a].split(" (")[0] for a in c["arms"]), c["billed_calls"],
             round(c["billed_cost"], 6), c["kept_runs"], round(c["kept_billed_cost"], 6), round(c["kept_cost"], 6),
             c["superseded_calls"], round(c["superseded_cost"], 6)] for c in crows]
    rows.append(["All models", "all three arms", ctot["billed_calls"], round(ctot["billed_cost"], 6),
                 ctot["kept_runs"], round(ctot["kept_billed_cost"], 6), round(ctot["kept_cost"], 6),
                 ctot["superseded_calls"], round(ctot["superseded_cost"], 6)])
    _, _, r = table(ws, r, ["Model", "Arms", "Requests billed", "Billed (US$)", "Runs kept",
                            "Kept, as the export bills them (US$)", "Kept, as the run records report them (US$)",
                            "Superseded requests", "Superseded (US$)"], rows,
                    [None, None, INT, USD4, INT, USD4, USD4, INT, USD4], bold_first=True)
    r = note(ws, r - 1, f"Source: {src['billing_name']} and work/*/by_model/*/run_records/*.json. Every request in the "
                        "export carries this study's application name. A model that appears in two arms is shown once, "
                        "with both arms named. Billed = kept as billed + superseded, exactly, in every row.")
    r = note(ws, r, "The last two kept columns cover the same 136 requests and differ only in reporting precision: the "
                    f"run records total US${ctot['kept_cost']:.6f} against US${ctot['kept_billed_cost']:.6f} in the "
                    f"export, a gap of US${abs(ctot['cost_reporting_gap']):.6f} across all kept runs. Neither figure is "
                    "adjusted to match the other.")
    r += 1
    r = h2(ws, r, "Table 2. Reconciliation")
    rows = [
        ["Requests billed in the provider's export", ctot["billed_calls"], round(ctot["billed_cost"], 6),
         "Every row of the export, all of it this study's extension application."],
        ["Runs kept in the analysis, as the export bills them", ctot["kept_runs"],
         round(ctot["kept_billed_cost"], 6),
         "One retained run record per scheduled run: 72 low-resource, 48 whole-document, 16 rule-discovery, matched to "
         "the export by generation identifier."],
        ["Superseded requests (billed, not kept)", ctot["superseded_calls"], round(ctot["superseded_cost"], 6),
         "Each was replaced by a later attempt of the same run. Kept plus superseded equals the billed total exactly."],
        ["Runs kept, as the run records report them", ctot["kept_runs"], round(ctot["kept_cost"], 6),
         "The provider's per-request cost as returned with each reply; this is the figure the arm sheets use."],
        ["Difference between the two kept figures", 0, round(ctot["cost_reporting_gap"], 6),
         "Reporting precision only, across 136 requests. The same requests appear in both routes."],
        ["Kept generation identifiers found in the export", ctot["kept_gids"], None,
         f"{ctot['gids_missing_from_export']} kept generation identifiers are missing from the export."],
    ]
    _, _, r = table(ws, r, ["Line", "Requests", "US$", "What it covers"], rows,
                    [None, INT, USD4, None], bold_first=True)
    ws.column_dimensions["D"].width = 86
    r = note(ws, r - 1, "Every generation identifier recorded in a run record appears in the export, and every export "
                        "row is either a kept run or a superseded attempt, so no kept run is unbilled and no billed "
                        "request is unaccounted for.")
    if ctot["unmatched_slugs"]:
        r = note(ws, r, "Export model slugs with no kept run: " + "; ".join(ctot["unmatched_slugs"]))
    r += 1
    r = h2(ws, r, "Table 3. Superseded attempts and replies with no usable content")
    rows = [[c["model_id"], c["superseded_calls"], round(c["superseded_cost"], 6), c["superseded_truncated"],
             c["superseded_no_content"]] for c in crows if c["superseded_calls"]]
    rows.append(["All models", sum(c["superseded_calls"] for c in crows),
                 round(sum(c["superseded_cost"] for c in crows), 6),
                 sum(c["superseded_truncated"] for c in crows),
                 sum(c["superseded_no_content"] for c in crows)])
    _, _, r = table(ws, r, ["Model", "Superseded requests", "Superseded (US$)",
                            "Stopped at the output cap", "Returned no visible content"], rows,
                    [None, INT, USD4, INT, INT], bold_first=True)
    r = note(ws, r - 1, "Stopped at the output cap = the export records finish reason 'length'. Returned no visible "
                        "content = every output token was a reasoning token, so nothing was left to save; that is the "
                        "behaviour recorded as Kimi K3's empty replies in DEVIATIONS.md, D2-02.")
    r = note(ws, r, "Superseded replies that had text are kept beside the reply that replaced them as "
                    "<run_id>_raw_superseded_<timestamp>.txt: "
                    + ", ".join(f"{ARM_LABEL[a].split(' (')[0]} {MODEL_LABEL.get(l, l)} {n}"
                                for (a, l), n in sorted(src["superseded_files"].items()))
                    + ". An attempt that returned no content left no file, which is why the file count is lower than "
                      "the superseded request count.")
    r += 1
    r = h2(ws, r, "Table 4. Token totals")
    rows = [
        ["Kept runs (the analysis)", ctot["kept_runs"], ctot["kept_prompt"], ctot["kept_completion"], None],
        ["All billed requests", ctot["billed_calls"], ctot["billed_prompt"], ctot["billed_completion"],
         ctot["billed_reasoning"]],
    ]
    _, _, r = table(ws, r, ["Scope", "Requests", "Input tokens", "Output tokens",
                            "of which reasoning tokens"], rows, [None, INT, INT, INT, INT], bold_first=True)
    r = note(ws, r - 1, "Run records report input and output tokens but not the reasoning split, so the reasoning column "
                        "is left empty for kept runs rather than guessed. Output tokens include reasoning tokens.")
    r += 1
    r = h2(ws, r, "Chart data (values used by the chart below)")
    cd0 = chart_block(ws, r, ["Model", "Billed (US$)", "Kept (US$)"],
                      [[c["model_id"], round(c["billed_cost"], 4), round(c["kept_cost"], 4)] for c in crows])
    r += len(crows) + 2
    ch = bar(ws, ref(ws, 1, cd0 + 1, 1, cd0 + len(crows)), ref(ws, 2, cd0, 3, cd0 + len(crows)),
             f"Whole documents drove the bill: US${ctot['billed_cost']:.2f} billed, US${ctot['kept_cost']:.2f} kept",
             "US$", [GREY, BLUE], horizontal=True, numfmt='"$"0.00', width=24, height=9.5, gap=40)
    place(ws, ch, f"A{r}", f"{src['billing_name']} and the retained run records (Table 1 above).")

    # ---------------------------------------------------------------- Dashboard
    dashboard(ws_dash, src, arms, prim, ctot, day)

    # ---------------------------------------------------------------- Data dictionary
    data_dictionary(wb, src)

    # ---------------------------------------------------------------- README
    readme(ws_readme, src, day)

    # stamp the workbook with the collection date rather than the build time, so re-running on unchanged
    # sources produces a byte-identical file
    try:
        stamped = datetime.strptime(day, "%Y-%m-%d")
    except ValueError:
        stamped = datetime(2026, 1, 1)
    wb.properties.created = stamped
    wb.properties.modified = stamped
    for sh in wb.worksheets:
        sh.page_setup.orientation = "landscape"
        sh.page_setup.paperSize = 9
        sh.sheet_properties.pageSetUpPr.fitToPage = True
        sh.page_setup.fitToWidth = 1
        sh.page_setup.fitToHeight = 0
    out_path = Path(out_path).expanduser()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    pin_timestamps(out_path, stamped)
    return wb, arms, prim, crows, ctot


def pin_timestamps(path, stamp):
    """Replace the save-time stamps so an unchanged rebuild gives a byte-identical file.

    openpyxl writes the moment of saving into docProps/core.xml and into every zip entry. Both are replaced
    here by the collection date already shown on the README sheet, so the workbook can be hashed and
    compared between runs. Nothing a reader sees changes.
    """
    import re
    import zipfile
    path = Path(path)
    iso = stamp.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    dt = (stamp.year, stamp.month, stamp.day, stamp.hour, stamp.minute, stamp.second)
    tmp = path.with_name(path.name + ".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                              rb"\g<1>" + iso + rb"\g<2>", data)
            info = zipfile.ZipInfo(item.filename, date_time=dt)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = item.external_attr
            zout.writestr(info, data)
    tmp.replace(path)


def arm_sheet(wb, src, arm, rows_in, prim, name, title, subtitle):
    scheduled = src["scheduled"][arm]
    ws = sheet(wb, name, title,
               subtitle + " " + UNSCORED_LINE,
               [26, 24, 11, 11, 11, 11, 12, 13, 13, 15, 12, 12, 12, 12, 12, 12, 13, 12])
    r = h2(ws, 4, f"Table 1. Coverage and format, per configuration (denominator = {scheduled} scheduled cases)")
    rows = [[m["name"], m["model_id"], m["scheduled"], m["runs"], m["replies_saved"], m["parsed"], m["schema_ok"],
             m["format_failures"], m["transport_errors"]] for m in rows_in]
    _, _, r = table(ws, r, ["Configuration", "Model identifier", "Cases scheduled", "Runs with a retained record",
                            "Replies received and saved", "Parsed as one JSON object", "Valid against the schema",
                            "Format failures (unparseable or schema-invalid)", "Transport errors"], rows,
                    [None, None, INT, INT, INT, INT, INT, INT, INT], bold_first=True)
    r = note(ws, r - 1, f"Source: analysis/{arm}_per_run.csv and work/{arm}/by_model/*/run_records/*.json. A reply is "
                        "counted as received when a non-empty reply was saved for that run; a reply that arrived but "
                        "could not be parsed is a format failure, not a missing reply.")
    r += 1
    r = h2(ws, r, "Table 2. Quotations, matched against the evidence that configuration was given")
    rows = [[m["name"], m["answers_with_quotes"], m["quotes"], m["quotes_in_evidence"],
             m["quotes"] - m["quotes_in_evidence"], m["quotes_in_excerpt"]] for m in rows_in]
    _, _, r = table(ws, r, ["Configuration", "Answers offering at least one quotation", "Quotations offered",
                            "Found in the evidence supplied", "Not found in the evidence supplied",
                            "Also found in the supplied excerpt"], rows,
                    [None, INT, INT, INT, INT, INT], bold_first=True)
    if arm == "access":
        r = note(ws, r - 1, "The evidence supplied here is the same frozen excerpt used in the primary study, so the "
                            "last two columns agree by construction.")
    else:
        r = note(ws, r - 1, "The evidence supplied here is the attached document or bundle, so a passage quoted from "
                            "elsewhere in it is legitimate and is counted as found. The last column counts how many of "
                            "those quotations also fall inside the short supplied excerpt; it is context, not a target.")
    r = note(ws, r, "A matched quotation shows only that the text exists in what the model was given. It says nothing "
                    "about whether the passage supports the advice: that is a scoring question, and these answers are "
                    "not yet scored.")
    r += 1
    r = h2(ws, r, "Table 3. Release decisions under both rule sets, and what the models declared")
    rows = [[m["name"], m["runs"], m["released_v1"], m["routed_v1"], m["blocked_v1"],
             m["released_v2"], m["routed_v2"], m["blocked_v2"],
             m["released_v2"] - m["released_v1"], m["consequential"]] for m in rows_in]
    fmts = [None, INT, INT, INT, INT, INT, INT, INT, "+#,##0;-#,##0;0", INT]
    if prim:
        rows.append([None] * 10)
        for p in prim[:2]:
            rows.append([p["name"] + " - the primary study, different arm", p["answers"], p["released_v1"], p["routed_v1"],
                         p["blocked_v1"], p["released_v2"], p["routed_v2"], p["blocked_v2"],
                         p["released_v2"] - p["released_v1"], None])
    _, _, r = table(ws, r, ["Configuration", "Answers", "v1: released with a warning", "v1: routed to a person",
                            "v1: blocked", "v2: released with a warning", "v2: routed to a person", "v2: blocked",
                            "Change in releases, v1 to v2", "Answers declaring a grade, approval or penalty"],
                    rows, fmts, bold_first=True)
    if prim:
        r = note(ws, r - 1, "The last two rows are the primary study frontier configurations (Astra and Luna, OpenAI Codex), "
                            "shown for comparison only. They are a different arm: 96 answers over 24 cases with two "
                            "repetitions each, collected in the primary study, human-scored in part, and they returned valid "
                            "structure in 96 of 96 answers, so their format-failure count is 0. Their timings and costs "
                            f"are not in the extension files and are left out rather than estimated ({NOT_HERE}).")
        r = note(ws, r, "Decision counts for those two rows come from the same arm-A re-run shown on the Revised checks "
                        "sheet, so they are like for like with the extension columns beside them.")
    else:
        r = note(ws, r - 1, f"Source: analysis/{arm}_per_run.csv. Released with a warning, routed and blocked are the "
                            "three decisions the checks can reach; they sum to the answer count in each row.")
    r += 1
    r = h2(ws, r, "Table 4. Time, cost and collection incidents")
    rows = [[m["name"], m["runs"], m["median_seconds"], m["cost"],
             round(m["cost"] / m["runs"], 6) if m["runs"] else None,
             m["prompt_tokens"], m["completion_tokens"], m["retried"], m["raised_cap"], m["truncated_kept"],
             m["superseded_files"]] for m in rows_in]
    _, _, r = table(ws, r, ["Configuration", "Runs", "Median seconds per answer", "Cost of kept runs (US$)",
                            "Cost per answer (US$)", "Input tokens", "Output tokens",
                            "Runs needing more than one attempt", "Runs re-sent with the raised output cap",
                            "Kept replies stopped at the output cap", "Superseded replies kept on disk"], rows,
                    [None, INT, DEC1, USD4, USD4, INT, INT, INT, INT, INT, INT], bold_first=True)
    r = note(ws, r - 1, "Median seconds is the send-to-finish time of the kept attempt. Cost is the provider's own "
                        "figure for the kept request only; superseded attempts are billed separately and appear on the "
                        "Costs and provenance sheet. The output cap was 4,000 tokens at first and 16,000 on re-sent runs "
                        "(DEVIATIONS.md, D2-01).")
    if arm == "documents":
        kimi = next((m for m in rows_in if m["label"] == "kimik3"), None)
        stuck = [x for x in src["per_run"][arm] if x["label"] == "kimik3" and not as_bool(x["parsed"])]
        if kimi and stuck:
            r += 1
            r = h2(ws, r, "Table 5. The one case left without a usable answer")
            rows = [[stuck[0]["case_id"], MODEL_LABEL["kimik3"],
                     "Empty replies on the first attempts, then a re-sent reply that stopped at the raised output cap "
                     "and could not be parsed",
                     src["records"][(arm, "kimik3", stuck[0]["run_id"])].get("reply_chars"),
                     src["records"][(arm, "kimik3", stuck[0]["run_id"])].get("max_tokens"),
                     src["records"][(arm, "kimik3", stuck[0]["run_id"])].get("finish_reason"),
                     stuck[0]["decision_v1"], as_float(stuck[0]["cost_usd"])]]
            _, _, r = table(ws, r, ["Case", "Configuration", "What happened", "Characters in the saved reply",
                                    "Output cap on the kept attempt", "Finish reason", "Decision under the checks",
                                    "Cost of the kept attempt (US$)"], rows,
                            [None, None, None, INT, INT, None, None, USD4], bold_first=True)
            ws.column_dimensions["C"].width = 60
            r = note(ws, r - 1, f"Kimi K3 therefore has {kimi['parsed']} usable answers of {kimi['scheduled']} scheduled "
                                "cases. Its coverage is stated that way, as answers over cases scheduled, never as a "
                                "rate over answers received.")
            r = note(ws, r, "Reading note: the published summary counts this case as a reply not received, while the "
                            "retained run record shows a reply that arrived and was cut off at the cap. The workbook "
                            "reports it as a reply received and a format failure, and says so here rather than silently "
                            "choosing one of the two.")
    if arm == "discovery":
        r += 1
        r = h2(ws, r, "Cases in this arm")
        r = note(ws, r, "Eight cases, fixed before collection: " + ", ".join(src["cases"][arm]) + ". They span all "
                 "three workflow groups (W1 evidence and version, W2 delegation and referral, W3 calculation and "
                 "configuration).")
        r = note(ws, r, "This arm cannot be pooled with the primary study or with the whole-document arm: the prompt differs, so "
                        "the comparison is between arms, not within one. A bundle can run to tens of thousands of "
                        "tokens, so cost and time here are driven by document size as much as by the model.")
    if arm == "access":
        r += 1
        r = note(ws, r, "What this arm can establish: whether models a school could legally self-host return the "
                        "required structure, quote text that exists in the excerpt, and what the checks do with their "
                        "answers. What it cannot: offline operation, privacy, latency or hardware feasibility. Prompts "
                        "still left the country and reached third-party infrastructure, and the provider serves these "
                        "models at higher precision than a quantized local copy, so these figures are an upper bound on "
                        "what a self-hosting school could achieve.")
    return ws


def dashboard(ws, src, arms, prim, ctot, day):
    all_arms = [m for a in ARMS for m in arms[a]]
    runs = sum(m["runs"] for m in all_arms)
    replies = sum(m["replies_saved"] for m in all_arms)
    parsed = sum(m["parsed"] for m in all_arms)
    schema = sum(m["schema_ok"] for m in all_arms)
    quotes = sum(m["quotes"] for m in all_arms)
    quotes_ok = sum(m["quotes_in_evidence"] for m in all_arms)
    eff = {e["outcome"]: e for e in effect_rows(src, "v2")}
    answers = sum(p["answers"] for p in prim)
    reproduced = sum(p["reproduced"] for p in prim)
    recorded = sum(p["recorded"] for p in prim)
    TILES = [
        (f"{runs}", "new answers collected", f"3 arms, {len(all_arms)} configurations, collected {day}", NAVY),
        (f"{parsed} of {runs}", "parsed as one JSON object", f"{schema} of {runs} also valid against the schema", NAVY),
        (f"{quotes_ok} of {quotes}", "quotations found in the evidence given", "existence only; support for the advice is unscored", BLUE),
        ("0 of " + str(runs), "new answers scored by a person", "correctness in these arms is not yet measured", RED),
        (f"{answers}", "answers re-run under the revised checks", f"{reproduced} of {recorded} recorded decisions reproduced", NAVY),
        (f"{eff['acceptable']['released_by_new']} and {eff['confirmed_error']['withheld_by_new']}",
         "acceptable answers recovered, errors contained", f"of {eff['acceptable']['answers']} scored acceptable and "
         f"{eff['confirmed_error']['answers']} confirmed wrong", BLUE),
        (f"{eff['unscored']['answers']}", "re-run answers nobody has scored", "counted neither as recovered help nor as contained error", GOLD),
        (f"US${ctot['billed_cost']:.2f}", "billed for the extension", f"US${ctot['kept_cost']:.2f} on the {ctot['kept_runs']} runs kept", NAVY),
    ]
    for k, (big, lab, sub, colr) in enumerate(TILES):
        rr0 = 4 + (k // 4) * 5
        cc0 = 2 + (k % 4) * 4
        ws.merge_cells(start_row=rr0, start_column=cc0, end_row=rr0 + 1, end_column=cc0 + 2)
        ws.merge_cells(start_row=rr0 + 2, start_column=cc0, end_row=rr0 + 2, end_column=cc0 + 2)
        ws.merge_cells(start_row=rr0 + 3, start_column=cc0, end_row=rr0 + 3, end_column=cc0 + 2)
        for rr_ in range(rr0, rr0 + 4):
            for cc in range(cc0, cc0 + 3):
                ws.cell(row=rr_, column=cc).fill = TILEFILL
        c = ws.cell(row=rr0, column=cc0, value=big)
        c.font = Font(name=FONT, size=20, bold=True, color=colr)
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        c = ws.cell(row=rr0 + 2, column=cc0, value=lab)
        c.font = Font(name=FONT, size=10, bold=True, color=INK)
        c.alignment = Alignment(horizontal="left", indent=1)
        c = ws.cell(row=rr0 + 3, column=cc0, value=sub)
        c.font = Font(name=FONT, size=8, color=MUTED)
        ws.row_dimensions[rr0 + 3].height = 22
        c.alignment = Alignment(horizontal="left", indent=1, vertical="top", wrap_text=True)
    r = 15
    r = h2(ws, r, "What the extension adds, in plain sentences", 2)
    lines = [
        f"the primary study is closed and is not changed by anything here. the extension adds three arms that collected {runs} new "
        f"answers on {day}, and one re-analysis that re-ran the release checks over the {sum(p['answers'] for p in prim)} "
        "answers already collected.",
        "The revised checks (arm A) were specified and hashed before the run. Re-running the frozen rules reproduced "
        f"{reproduced} of the {recorded} decisions recorded at the September 18 score lock, so the before-and-after "
        f"comparison is like for like. They recovered {eff['acceptable']['released_by_new']} acceptable answers that the "
        f"frozen checks had withheld, contained {eff['confirmed_error']['withheld_by_new']} of "
        f"{eff['confirmed_error']['answers']} answers confirmed wrong by hand, and newly withheld "
        f"{eff['acceptable']['withheld_by_new']} answers scored acceptable.",
        "Three open-weight models a school could legally self-host answered the same 24 cases (arm G). Two models "
        "answered all 24 cases with the whole governing document attached (arm C). Two answered eight cases with every "
        "retained document of the institution attached and no passage pre-selected (arm D).",
        f"Of the {runs} new answers, {replies} arrived and were saved, {parsed} parsed as one JSON object and {schema} "
        f"were valid against the required schema. The models offered {quotes} quotations, of which {quotes_ok} were "
        "found word for word in the evidence they had been given.",
        f"the extension was billed US${ctot['billed_cost']:.4f} over {ctot['billed_calls']} requests. The "
        f"{ctot['kept_runs']} runs the analysis keeps account for US${ctot['kept_cost']:.4f} of that; a further "
        f"US${ctot['superseded_cost']:.4f} over {ctot['superseded_calls']} requests paid for attempts that were "
        "superseded when a reply came back empty or stopped at the output cap. Superseded spending is reported, "
        "never netted off.",
    ]
    for t in lines:
        c = ws.cell(row=r, column=2, value=t)
        c.font = F_BODY
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=2, end_row=r + 1, end_column=16)
        ws.row_dimensions[r].height = 15
        ws.row_dimensions[r + 1].height = 15
        r += 3
    c = ws.cell(row=r, column=2, value="Correctness for the new arms is not yet scored by a person.")
    c.font = Font(name=FONT, size=11, bold=True, color=RED)
    r += 1
    c = ws.cell(row=r, column=2, value=UNSCORED_LINE + " A configuration that is blocked more often may be safer or "
                                      "merely messier in its formatting, and only human scoring can separate the two. "
                                      "Arm B of the extension protocol does that scoring; until it lands, no ranking, "
                                      "rate or quality claim should be read off these sheets.")
    c.font = F_BODY
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r, start_column=2, end_row=r + 1, end_column=16)
    r += 3
    note(ws, r, "Tiles and sentences are computed by the build script from the files named on each sheet. Every count "
                "carries its denominator; no rate is given without one.", 2)


def data_dictionary(wb, src):
    ws = sheet(wb, "Data dictionary", "Data dictionary",
               "Every column and measure used in this workbook, and what none of them establishes.",
               [34, 96, 30])
    r = h2(ws, 4, "Columns and measures")
    rows = [
        ("Arm", "A named part of the extension. Arm A re-runs the checks over answers already collected; arms C, D and G "
                "collect new answers. Arms are reported separately and are never pooled.", "All sheets"),
        ("Configuration", "One model under one arm, named by the public label used in the extension results files "
                          "(for example Gemma 3 12B, Kimi K3) or, for primary comparison rows, by the study label "
                          "(Astra, Luna, Fable 5.1, Haiku 4.5).", "All sheets"),
        ("Cases scheduled", "The runs fixed in the arm's manifest before collection: 24 cases for the low-resource "
                            "and whole-document arms, 8 for rule discovery.", "Arm sheets"),
        ("Replies received and saved", "A non-empty reply came back for that run and was written to disk. A reply that "
                                       "arrived and could not be parsed is still a reply received.", "Arm sheets"),
        ("Parsed as one JSON object", "The frozen parser read the saved reply as a single JSON object. One enclosing "
                                      "code fence may be removed; nothing else is repaired.", "Arm sheets"),
        ("Schema ok", "The parsed object carried every field the study's schema requires, each of the required type. "
                      "Every schema failure in the extension was the same fault: a numeric_results entry with an empty or "
                      "non-text field, usually a missing unit. Schema ok says nothing about whether the content is "
                      "right.", "Arm sheets, Data dictionary"),
        ("Format failures", "Runs that were unparseable or schema-invalid: the complement of schema ok over the runs "
                            "with a retained record.", "Arm sheets"),
        ("Quotations offered", "Entries in the answer's citations list.", "Arm sheets, Dashboard"),
        ("Quotation matched against evidence", "The quoted string, after normalization, appears word for word inside "
                                               "the text the model was given: the frozen excerpt in the low-resource "
                                               "arm, the attached document in the whole-document arm, the whole bundle "
                                               "in rule discovery. It establishes that the passage exists in what the "
                                               "model was shown. It does not establish that the passage is the "
                                               "governing rule, that it supports the advice, or that the answer is "
                                               "right.", "Arm sheets, Dashboard"),
        ("Released (released with a warning)", "The checks found no failure and no declared consequential action, so "
                                               "the answer goes to the user with a standard caution that the checks "
                                               "are partial. It is not a statement that the answer is correct.",
         "Arm sheets, Revised checks"),
        ("Routed", "The answer declared a consequential action (record a grade, approve a request, penalize a "
                   "student), so it is sent to a person instead of released. Routing is a policy choice about who "
                   "decides, not a finding that the answer is wrong.", "Arm sheets, Revised checks"),
        ("Blocked", "A check failed (format, quotation membership, or a recomputation), so the answer is withheld. "
                    "Blocking takes precedence over routing. A blocked answer may still have been correct.",
         "Arm sheets, Revised checks"),
        ("Withheld", "Routed or blocked: the answer did not reach the user unaided.", "Revised checks"),
        ("Rule set v1 / v2 / v2b", "v1 = the four release checks frozen on 15 September 2026. v2 = the revised checks "
                                   "specified and hashed before the re-run. v2b = an exploratory variant that also "
                                   "reads decisions and numbers out of the answer text; reported separately and never "
                                   "merged into v2.", "Revised checks, arm sheets"),
        ("Outcome class", "The primary study human verdict on a re-run answer, unchanged: scored acceptable, scored serious "
                          "(Major or Critical), confirmed wrong by hand, or not scored by a person.", "Revised checks"),
        ("Recovered / newly withheld", "Recovered = withheld under v1 and released under the new rules. Newly withheld "
                                       "= released under v1 and withheld under the new rules.", "Revised checks"),
        ("Median seconds per answer", "Send-to-finish time of the kept attempt for that run, median over the "
                                      "configuration's runs.", "Arm sheets"),
        ("Cost of kept runs", "The provider's own cost figure for the kept request of each run, summed. Superseded "
                              "attempts are billed separately and are shown on the Costs and provenance sheet.",
         "Arm sheets, Costs and provenance"),
        ("Billed", "Every request charged in the provider's activity export for this study's extension application. "
                   "The export and the run records report the same requests to different precision, so kept cost is "
                   "shown both ways and neither is adjusted to match the other.", "Costs and provenance"),
        ("Kept", "The latest attempt of each scheduled run: the request the analysis uses.", "Costs and provenance"),
        ("Superseded", "A billed request replaced by a later attempt of the same run, because the reply was empty or "
                       "stopped at the output cap. Kept in the record, never deleted, never netted off the bill.",
         "Costs and provenance"),
        ("Stopped at the output cap", "The provider recorded finish reason 'length': the model was still writing when "
                                      "the output limit was reached. A cut-off reply is a collection setting's "
                                      "failure, not a failure of the model's ability to return the format.",
         "Costs and provenance, arm sheets"),
        ("Returned no visible content", "Every output token billed for that request was a reasoning token, so no "
                                        "answer text came back.", "Costs and provenance"),
        ("Input / output tokens", "Prompt and completion tokens as the provider reports them. Output tokens include "
                                  "reasoning tokens where a model emits them.", "Arm sheets, Costs and provenance"),
    ]
    _, _, r = table(ws, r, ["Field", "Definition", "Sheets"], [list(x) for x in rows], [None, None, None],
                    bold_first=True)
    r += 1
    r = h2(ws, r, "What none of these measures establishes")
    for t in [
        "Correctness. No answer from arms C, D or G has been scored by a person. Parsing, schema validity, quotation "
        "membership and check decisions are all properties of the form of an answer and of the evidence it cites; none "
        "of them shows that the advice is right.",
        "A ranking of models or providers. The arms differ in prompt, evidence and case count, and are reported "
        "separately for that reason. A configuration blocked more often may be safer or merely messier.",
        "A rate for any population. These are counts over 24 cases (8 in rule discovery) built from nine Philippine "
        "universities' published rules. They describe these cases and these configurations only.",
        "Anything about offline or self-hosted operation. Every request in the extension went through a hosted aggregator, "
        "so no privacy, latency, sovereignty or hardware claim follows from the low-resource arm.",
        "That the revised checks would behave this way on new answers. Every answer in arm A was collected under the "
        "frozen checks; no model knew the rules had changed.",
        "That a matched quotation is the governing rule. In the whole-document and rule-discovery arms the model chose "
        "its own passage, and whether it chose the right one is a scoring question.",
    ]:
        c = ws.cell(row=r, column=1, value=t)
        c.font = F_BODY
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=1, end_row=r + 1, end_column=3)
        r += 3
    return ws


def readme(ws, src, day):
    rrec = src["checks_run_record"]
    lines = [
        ("Purpose", "The extension to the assessment-control capstone study, in one workbook: three new "
                    "collection arms and one re-analysis of answers already collected. It sits beside "
                    "Analysis_Workbook_FINAL.xlsx and Analysis_Workbook_Supplement_v2.xlsx, which are unchanged by "
                    "anything here."),
        ("the primary study is closed", "Its cases, checks, answers, scores and results are not modified. Primary-study figures that "
                              "appear here are re-computed from the re-run of the checks over the same "
                              "answers, and are labelled as a different arm wherever they sit beside extension rows."),
        ("Arm A, revised checks", "The four frozen release checks were revised to the recommendation the paper itself "
                                  "makes in Section 6.3, specified and hashed before the run, and applied to the 192 "
                                  "answers already collected. No new answers, no paid calls, no score revised."),
        ("Arm G, low-resource models", "Three open-weight models in the 8B-30B range answered the same 24 cases with "
                                       "the same supplied excerpt, through a hosted aggregator. Not an offline "
                                       "deployment."),
        ("Arm C, whole documents", "Two models answered all 24 cases with the entire governing document attached and "
                                   "no passage indicated."),
        ("Arm D, rule discovery", "Two models answered eight pre-selected cases with every retained document of the "
                                  "institution attached and no passage indicated."),
        ("", ""),
        ("Correctness", "Not scored for arms C, D and G. No person has scored a single answer from those arms, so "
                        "nothing on those sheets is a correctness result. Arm B of the extension protocol does that "
                        "scoring."),
        ("How to read", "Every count carries its denominator, and no rate appears without one. Tables sit on each "
                        "sheet; charts sit below a 'Chart data' block holding the exact values plotted. Grey = the "
                        "frozen rules or the billed total (baseline); blue = the revised rules or the kept total; gold "
                        "= unscored; red = a limit on what can be claimed. A source note sits under every table and "
                        "every chart."),
        ("Privacy", "No raw model answer appears in this workbook. Configurations are named only by the labels already "
                    "public in the extension results files, and answers are identified by case, configuration and "
                    "repetition."),
        ("", ""),
        ("Caveats", ""),
        ("Unscored arms", "Coverage, format, quotation membership and check decisions describe what the models "
                          "produced, not whether they were right."),
        ("Hindsight in arm A", "The revised rules were written after the author had seen the v1 results. Each follows "
                               "the recommendation published on 18 September 2026, before this arm existed, and the "
                               "hashes below fix what was claimed before the data was touched; the risk cannot be "
                               "removed by assertion."),
        ("Thin containment evidence", "Two answers confirmed wrong by hand is a thin basis for the containment claim. "
                                      "Report the count, never a rate."),
        ("Arms are not comparable", "Prompt, evidence and case count differ between arms. Figures are reported per "
                                    "arm and are not pooled."),
        ("Collection incidents", "The output cap was raised after the first run and affected replies were re-sent; one "
                                 "whole-document case has no usable answer. Both are recorded on the sheets and in "
                                 "DEVIATIONS.md (D2-01, D2-02)."),
        ("", ""),
        ("Sheets and sources", ""),
        ("Dashboard", "Tiles and plain sentences (sources as on the linked sheets)."),
        ("Revised checks", "analysis/checks_v2_per_answer.csv (192 answers); analysis/checks_v2_changes.csv; "
                           "08_RESULTS_PUBLIC/checks_v2_run_record.json."),
        ("Low-resource models", "analysis/access_per_run.csv; work/access/MANIFEST.json; "
                                "work/access/by_model/*/run_records/*.json."),
        ("Whole documents", "analysis/documents_per_run.csv; work/documents/MANIFEST.json; "
                            "work/documents/by_model/*/run_records and raw_outputs."),
        ("Rule discovery", "analysis/discovery_per_run.csv; work/discovery/MANIFEST.json; "
                           "work/discovery/by_model/*/run_records/*.json."),
        ("Costs and provenance", src["billing_name"] + " (the provider's own export) against "
                                 "work/*/by_model/*/run_records/*.json."),
        ("Data dictionary", "Field definitions, and what none of these measures establishes."),
        ("", ""),
    ]
    if rrec:
        lines.append(("Arm A freeze record", f"Specification SHA-256 {rrec.get('spec_sha256', NOT_HERE)}; "
                                             f"implementation SHA-256 {rrec.get('implementation_sha256', NOT_HERE)}; "
                                             f"{rrec.get('answers_scanned', NOT_HERE)} answers scanned, "
                                             f"{len(rrec.get('v1_reproduction_mismatches', []))} reproduction "
                                             "mismatches."))
    lines += [
        ("Answers collected", day + " (latest finish time in the retained run records; the workbook carries no build "
                                    "timestamp, so re-running it on unchanged sources gives an identical file)."),
        ("Rebuild", "python3 -B code/extensions/round2/build_extension_workbook.py --package <30_ROUND2_EXTENSION> "
                    "--out <package>/analysis/Analysis_Workbook_Extension.xlsx  (reads sources only; rewrites this "
                    "workbook)."),
    ]
    rr = 4
    for k, v in lines:
        a = ws.cell(row=rr, column=1, value=k)
        a.font = F_BOLD if k and not v else F_BODY
        a.alignment = Alignment(vertical="top")
        b = ws.cell(row=rr, column=2, value=v)
        b.font = F_BODY
        b.alignment = Alignment(wrap_text=True, vertical="top")
        rr += 1
    return ws


# ----------------------------------------------------------------------------------------------
# 5. MAIN
# ----------------------------------------------------------------------------------------------
def main(argv=None):
    a = parse_args(argv)
    src = load(a.package, a.billing)
    wb, arms, prim, crows, ctot = build(src, a.out)
    out = Path(a.out).expanduser()
    print(f"saved {out}")
    print(f"  sheets: {', '.join(wb.sheetnames)}")
    total_runs = sum(m["runs"] for arm in ARMS for m in arms[arm])
    print(f"  new answers collected: {total_runs}; parsed "
          f"{sum(m['parsed'] for arm in ARMS for m in arms[arm])}; schema ok "
          f"{sum(m['schema_ok'] for arm in ARMS for m in arms[arm])}")
    print(f"  answers re-run under the revised checks: {sum(p['answers'] for p in prim)}; "
          f"recorded decisions reproduced {sum(p['reproduced'] for p in prim)}/{sum(p['recorded'] for p in prim)}")
    print(f"  billed US${ctot['billed_cost']:.4f} over {ctot['billed_calls']} requests; "
          f"kept US${ctot['kept_cost']:.4f} over {ctot['kept_runs']} runs")
    print("  correctness for arms C, D and G is not scored; these are coverage, format, quotation and decision counts.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
